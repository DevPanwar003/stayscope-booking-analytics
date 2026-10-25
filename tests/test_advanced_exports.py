"""Read generated workbooks back to verify values, scope and safe cells."""
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.exports import advanced_report_tables, excel_bytes


def test_advanced_workbook_reconciles_each_dimension(clean):
    original = clean.copy(deep=True)
    book = pd.read_excel(BytesIO(excel_bytes(clean, {'currency': 'USD', 'minimum_group_size': 3, 'synthetic': True})), sheet_name=None)
    assert list(book) == ['Summary', 'Filtered bookings', 'Cancellation benchmarks', 'Value bands', 'Report guide']
    benchmarks = book['Cancellation benchmarks']
    assert set(benchmarks.dimension) == {'channel', 'room_type', 'city', 'property', 'star_rating', 'booking_period'}
    for _, frame in benchmarks.groupby('dimension'):
        assert frame.bookings.sum() == 4
        assert frame.cancelled.sum() == 1
        assert frame.cancelled_booking_value.sum() == 200
        assert frame.overall_cancellation_rate.eq(.25).all()
        assert frame.currency.eq('USD').all()
        assert frame.minimum_group_size.eq(3).all()
    ota = benchmarks[(benchmarks.dimension == 'channel') & (benchmarks.segment == 'OTA')].iloc[0]
    assert ota.cancellation_rate == .5 and ota.difference_pp == 25
    assert ota.small_sample and ota.ci_lower < .5 < ota.ci_upper
    assert book['Value bands'].bookings.sum() == 4
    assert book['Filtered bookings'].booking_id.tolist() == clean.booking_id.tolist()
    assert 'not real hotel results' in book['Report guide'].to_string()
    pd.testing.assert_frame_equal(clean, original)


def test_filtered_workbook_excludes_other_rows_and_uses_threshold(clean):
    data = clean[clean.channel == 'Direct']
    tables = advanced_report_tables(data, {'minimum_group_size': 1})
    groups = tables['Cancellation benchmarks']
    assert groups.overall_cancellation_rate.eq(0).all()
    assert groups.cancelled_booking_value.eq(0).all()
    assert not groups.small_sample.any()
    assert groups.currency.eq('Unspecified units').all()


def test_empty_workbook_preserves_all_sheets_and_headers(clean):
    book = pd.read_excel(BytesIO(excel_bytes(clean.iloc[:0])), sheet_name=None)
    assert book['Cancellation benchmarks'].empty
    assert {'dimension', 'segment', 'bookings', 'ci_lower', 'currency'} <= set(book['Cancellation benchmarks'])
    assert book['Value bands'].empty
    assert book['Filtered bookings'].empty
    assert not book['Report guide'].empty


@pytest.mark.parametrize('size', [4, 24])
def test_missing_optional_fields_are_safe_without_changing_original_columns(clean, size):
    data = pd.concat([clean] * (size // 4), ignore_index=True)[['booking_id', 'booking_date', 'status', 'booking_value']]
    book = pd.read_excel(BytesIO(excel_bytes(data)), sheet_name=None)
    assert list(book['Filtered bookings']) == list(data)
    for dimension in ['property', 'city', 'channel', 'room_type', 'star_rating']:
        rows = book['Cancellation benchmarks'].query('dimension == @dimension')
        assert rows.segment.tolist() == ['Unknown']
        assert rows.bookings.sum() == size


def test_new_sheets_are_spreadsheet_safe_and_formatted(clean):
    clean.loc[0, 'channel'] = '=1+1'
    book = load_workbook(BytesIO(excel_bytes(clean, {'currency': '=CMD()', 'synthetic': True})))
    sheet = book['Cancellation benchmarks']
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.freeze_panes == 'A2'
    assert sheet.cell(2, headers['cancellation_rate']).number_format == '0.0%'
    assert sheet.cell(2, headers['booking_value']).number_format == '#,##0.00'
    for worksheet in book:
        assert all(cell.data_type != 'f' for row in worksheet for cell in row)
    assert any(cell.value == "'=1+1" for row in sheet for cell in row)


def test_unavailable_cancelled_values_not_exported_as_zero(clean):
    clean.loc[1, 'booking_value'] = float('nan')
    frame = pd.read_excel(BytesIO(excel_bytes(clean)), sheet_name='Cancellation benchmarks')
    ota = frame[(frame.dimension == 'channel') & (frame.segment == 'OTA')].iloc[0]
    assert ota.cancelled_booking_value == 'Unavailable'
    assert ota.invalid_cancelled_values == 1
    assert ota.cancellation_rate == .5
