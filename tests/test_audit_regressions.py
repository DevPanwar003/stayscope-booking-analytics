"""Regression coverage added during the quality audit; uses disposable inputs."""
from io import BytesIO
import sqlite3
import pandas as pd
import pytest
from src.database import initialize, save_dataset, quality_report, datasets, reset_demo, load_dataset
from src.exports import excel_bytes
from src.validation import DataError, read_upload, validate


def test_duplicate_csv_headers_rejected():
    content = b'booking_id,booking_date,status,booking_value,booking_value\nA,2025-01-01,Confirmed,100,999\n'
    with pytest.raises(DataError, match='same field|Duplicate'):
        validate(read_upload(content, 'duplicate.csv'))


def test_literal_na_identifier_is_preserved():
    content = b'booking_id,booking_date,status,booking_value\nNA,2025-01-01,Confirmed,100\n'
    assert validate(read_upload(content, 'bookings.csv')).accepted.booking_id.tolist() == ['NA']


def test_blank_csv_records_are_not_silently_skipped():
    content = b'booking_id,booking_date,status,booking_value\n\nA,2025-01-01,Confirmed,100\n'
    result = validate(read_upload(content, 'bookings.csv'))
    assert result.summary['input_rows'] == 2
    assert result.accepted.source_row.tolist() == [3]


def test_audit_retains_leading_zero_ids(tmp_path, raw):
    raw['booking_id'] = ['001', '002', '003', '004']
    raw.loc[0, 'booking_value'] = -1
    db = tmp_path / 'audit.db'
    initialize(db)
    dataset_id, _ = save_dataset(db, validate(raw), 'audit')
    assert quality_report(db, dataset_id)[2].booking_id.tolist() == ['001']


def test_excel_handles_control_characters_without_crashing(clean):
    clean.loc[0, 'city'] = 'Goa\x01test'
    output = excel_bytes(clean)
    assert pd.read_excel(BytesIO(output), sheet_name='Filtered bookings').loc[0, 'city'] == 'Goa\\x01test'


def test_transaction_rolls_back_partial_import(tmp_path, raw):
    result = validate(raw)
    result.accepted.loc[1, 'booking_id'] = result.accepted.loc[0, 'booking_id']
    db = tmp_path / 'atomic.db'
    initialize(db)
    with pytest.raises(sqlite3.IntegrityError):
        save_dataset(db, result, 'must roll back')
    assert datasets(db).empty
    with sqlite3.connect(db) as connection:
        assert connection.execute('SELECT COUNT(*) FROM bookings').fetchone()[0] == 0


def test_sql_values_are_not_executed(tmp_path, raw):
    value = "x'); DROP TABLE bookings; --"
    raw.loc[0, 'booking_id'] = value
    db = tmp_path / 'sql.db'
    initialize(db)
    dataset_id, _ = save_dataset(db, validate(raw), value)
    with sqlite3.connect(db) as connection:
        assert connection.execute('SELECT booking_id FROM bookings WHERE dataset_id=? AND booking_id=?',
                                  (dataset_id, value)).fetchone()[0] == value
        assert connection.execute('SELECT COUNT(*) FROM bookings').fetchone()[0] == 4


def test_atomic_demo_restore_rolls_back_on_insert_failure(tmp_path, raw):
    db = tmp_path / 'reset.db'
    initialize(db)
    demo_id, _ = save_dataset(db, validate(raw), 'Demo', is_demo=True)
    user_id, _ = save_dataset(db, validate(raw), 'Private upload')
    broken = validate(raw)
    broken.accepted.loc[1, 'booking_id'] = 'A'
    with pytest.raises(sqlite3.IntegrityError):
        reset_demo(db, 'RESET DEMO', replacement=broken)
    assert len(load_dataset(db, demo_id)) == 4
    assert len(load_dataset(db, user_id)) == 4
    reset_demo(db, 'RESET DEMO', replacement=validate(raw))
    assert len(datasets(db)) == 2
    assert len(load_dataset(db, user_id)) == 4


def test_simultaneous_equivalent_imports(tmp_path, raw):
    from concurrent.futures import ThreadPoolExecutor
    db = tmp_path / 'concurrent.db'
    initialize(db)
    result = validate(raw)
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: save_dataset(db, result, 'same'), range(2)))
    assert sum(created for _, created in results) == 1
    assert len({identifier for identifier, _ in results}) == 1


def test_duplicate_excel_headers_rejected():
    buffer = BytesIO()
    pd.DataFrame([['A', '2025-01-01', 'Confirmed', 100, 999]], columns=[
        'booking_id', 'booking_date', 'status', 'booking_value', 'booking_value']).to_excel(buffer, index=False)
    with pytest.raises(DataError, match='Duplicate'):
        read_upload(buffer.getvalue(), 'bad.xlsx')


def test_csv_unequal_fields_are_not_shifted_into_index():
    content = b'booking_id,booking_date,status,booking_value\nA,2025-01-01,Confirmed,100,999\n'
    with pytest.raises(DataError, match='same number'):
        read_upload(content, 'wide.csv')


def test_excel_literal_na_id(raw):
    raw.loc[0, 'booking_id'] = 'NA'
    output = BytesIO()
    raw.to_excel(output, index=False)
    assert validate(read_upload(output.getvalue(), 'id.xlsx')).accepted.booking_id.iloc[0] == 'NA'


def test_timezone_normalizes_to_utc_day(raw):
    raw.loc[0, 'booking_date'] = '2025-01-01T00:30:00+05:30'
    assert validate(raw).accepted.booking_date.iloc[0] == pd.Timestamp('2024-12-31')


def test_no_retained_and_negative_markup(clean):
    from src.analytics import kpis
    assert kpis(clean[clean.status == 'Cancelled'])['Retained markup'] is None
    only = clean.iloc[[0]].copy()
    only['cost'] = 150
    assert kpis(only)['Retained markup'] == -50


def test_empty_exports_preserve_schema(clean):
    from src.exports import csv_bytes, summary_frame
    empty = clean.iloc[:0]
    assert pd.read_csv(BytesIO(csv_bytes(empty))).empty
    workbook = pd.read_excel(BytesIO(excel_bytes(empty)), sheet_name=None)
    assert workbook['Filtered bookings'].empty
    assert list(workbook['Filtered bookings'].columns) == list(clean.columns)
    assert summary_frame(empty).query('name == "Total bookings"').value.iloc[0] == 0


def test_formula_neutralization_in_excel(clean):
    from openpyxl import load_workbook
    clean.loc[0, 'city'] = '  =1+1'
    workbook = load_workbook(BytesIO(excel_bytes(clean)))
    sheet = workbook['Filtered bookings']
    city_column = list(clean.columns).index('city') + 1
    cell = sheet.cell(2, city_column)
    assert cell.data_type == 's'
    assert cell.value.startswith("'")


def test_different_quarantined_records_do_not_reuse_an_unrelated_audit(tmp_path, raw):
    db = tmp_path / 'snapshots.db'
    initialize(db)
    raw.loc[0, 'booking_value'] = -1
    first, _ = save_dataset(db, validate(raw), 'first')
    raw.loc[0, 'booking_id'] = 'DIFFERENT-REJECTED-BOOKING'
    second, created = save_dataset(db, validate(raw), 'second')
    assert created and second != first
    assert quality_report(db, second)[2].booking_id.tolist() == ['DIFFERENT-REJECTED-BOOKING']
    assert save_dataset(db, validate(raw), 'second-again') == (second, False)
